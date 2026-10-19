
"""Extract ordered text, tables, and images from PowerPoint presentations."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pdf_processor import DocumentElement
from settings import Settings
from visual_processor import analyze_visual


def process_pptx(
    path: Path,
    settings: Settings | None = None,
) -> list[DocumentElement]:
    """Extract slide content while preserving slide provenance."""

    path = path.resolve()

    if not path.is_file():
        raise ValueError(f"PPTX file does not exist: {path}")

    try:
        presentation = Presentation(path)
    except Exception as error:
        raise ValueError("PPTX file is unreadable or invalid.") from error

    elements: list[DocumentElement] = []
    image_cache = {}

    def append_element(
        content_type: str,
        content: str,
        slide_number: int,
        extraction_method: str,
        extraction_model: str | None = None,
    ) -> None:
        if not content.strip():
            return

        elements.append(
            DocumentElement(
                content_type=content_type,
                content=content.strip(),
                page_number=slide_number,
                element_index=len(elements) + 1,
                extraction_method=extraction_method,
                extraction_model=extraction_model,
            )
        )

    for slide_number, slide in enumerate(presentation.slides, start=1):
        # Sort by vertical position, then horizontal position.
        # This approximates the order a person reads slide content.
        shapes = sorted(
            slide.shapes,
            key=lambda shape: (shape.top, shape.left),
        )

        for shape in shapes:
            if shape.has_table:
                rows = []

                for row in shape.table.rows:
                    cells = [
                        cell.text.strip().replace("\n", " ")
                        for cell in row.cells
                    ]
                    rows.append(" | ".join(cells))

                append_element(
                    "table",
                    "\n".join(rows),
                    slide_number,
                    "python_pptx_table",
                )

            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                if settings is None:
                    continue

                image = shape.image
                image_format = image.ext.lower()

                if image_format == "jpg":
                    image_format = "jpeg"

                if image_format not in {"png", "jpeg"}:
                    continue

                image_bytes = image.blob

                # Reuse VLM results when the same image appears again.
                image_key = image.sha1

                if image_key not in image_cache:
                    try:
                        image_cache[image_key] = analyze_visual(
                            image_bytes,
                            image_format,
                            settings,
                        )
                    except Exception as error:
                        print(
                            "Warning: PPTX image analysis failed "
                            f"on slide {slide_number}: {error}"
                        )
                        continue

                analysis = image_cache[image_key]

                append_element(
                    analysis.content_type,
                    analysis.content,
                    slide_number,
                    f"{settings.vlm_provider}_vision",
                    analysis.model,
                )

            elif shape.has_text_frame:
                text = shape.text.strip()

                append_element(
                    "text",
                    text,
                    slide_number,
                    "python_pptx_text",
                )

    return elements
