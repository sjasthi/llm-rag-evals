
"""Tests for structured DOCX text and table extraction."""

import sys
import tempfile
import unittest
from pathlib import Path
from io import BytesIO
from unittest.mock import patch

from docx import Document
from docx.shared import Inches
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "rag"))

from docx_processor import process_docx
from visual_processor import VisualAnalysis

class TestDocxProcessor(unittest.TestCase):

    def test_preserves_text_and_table_order(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test_document.docx"

            document = Document()

            document.add_heading("Student Services Guide", level=1)

            document.add_paragraph(
                "Students can access technology support."
            )

            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Service"
            table.cell(0, 1).text = "Location"
            table.cell(1, 0).text = "IT Help Desk"
            table.cell(1, 1).text = "Library"

            document.add_paragraph(
                "Support is available Monday through Friday."
            )

            document.save(path)

            elements = process_docx(path)

            self.assertEqual(
                [element.content_type for element in elements],
                ["text", "text", "table", "text"],
            )

            self.assertEqual(
                [element.element_index for element in elements],
                [1, 2, 3, 4],
            )

            self.assertIn(
                "IT Help Desk | Library",
                elements[2].content,
            )

            self.assertEqual(
                elements[2].extraction_method,
                "python_docx_table",
            )

            self.assertTrue(
                all(element.page_number is None for element in elements)
            )

            self.assertTrue(
                all(element.extraction_model is None for element in elements)
            )

    def test_extracts_inline_image_in_document_order(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "image_document.docx"

            # Generate a small PNG image entirely in memory.
            image_buffer = BytesIO()
            Image.new("RGB", (100, 60), "white").save(
                image_buffer,
                format="PNG",
            )
            image_buffer.seek(0)

            document = Document()
            paragraph = document.add_paragraph()
            paragraph.add_run("Before the image. ")

            paragraph.add_run().add_picture(
                image_buffer,
                width=Inches(1),
            )

            paragraph.add_run(" After the image.")
            document.save(path)

            mock_settings = type(
                "MockSettings",
                (),
                {"vlm_provider": "ollama"},
            )()

            fake_analysis = VisualAnalysis(
                content_type="diagram",
                content="The workflow proceeds from Start to Finish.",
                model="test-vision-model",
            )

            with patch(
                "docx_processor.analyze_visual",
                return_value=fake_analysis,
            ) as mock_analyze:
                elements = process_docx(path, settings=mock_settings)

            self.assertEqual(
                [element.content_type for element in elements],
                ["text", "diagram", "text"],
            )

            self.assertEqual(
                [element.element_index for element in elements],
                [1, 2, 3],
            )

            self.assertEqual(elements[0].content, "Before the image.")
            self.assertEqual(
                elements[1].content,
                "The workflow proceeds from Start to Finish.",
            )
            self.assertEqual(elements[2].content, "After the image.")

            self.assertEqual(
                elements[1].extraction_method,
                "ollama_vision",
            )
            self.assertEqual(
                elements[1].extraction_model,
                "test-vision-model",
            )

            mock_analyze.assert_called_once()
            args = mock_analyze.call_args.args

            self.assertTrue(args[0].startswith(b"\x89PNG"))
            self.assertEqual(args[1], "png")

    def test_reuses_visual_analysis_for_repeated_image(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "repeated_image.docx"

            image_path = Path(directory) / "workflow.png"
            Image.new("RGB", (100, 60), "white").save(image_path)

            document = Document()

            document.add_paragraph("First workflow diagram:")
            document.add_paragraph().add_run().add_picture(
                str(image_path),
                width=Inches(1),
            )

            document.add_paragraph("Second workflow diagram:")
            document.add_paragraph().add_run().add_picture(
                str(image_path),
                width=Inches(1),
            )

            document.save(path)

            mock_settings = type(
                "MockSettings",
                (),
                {"vlm_provider": "ollama"},
            )()

            fake_analysis = VisualAnalysis(
                content_type="diagram",
                content="Start leads to Review, then Approval.",
                model="test-vision-model",
            )

            with patch(
                "docx_processor.analyze_visual",
                return_value=fake_analysis,
            ) as mock_analyze:
                elements = process_docx(path, settings=mock_settings)

            self.assertEqual(
                [element.content_type for element in elements],
                ["text", "diagram", "text", "diagram"],
            )

            self.assertEqual(
                [element.element_index for element in elements],
                [1, 2, 3, 4],
            )

            self.assertEqual(elements[1].content, elements[3].content)

            # The same image should be analyzed only once.
            mock_analyze.assert_called_once()

if __name__ == "__main__":
    unittest.main()
