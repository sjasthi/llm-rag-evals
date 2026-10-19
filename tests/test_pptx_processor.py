
"""Unit tests for multimodal PowerPoint extraction."""

import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from pptx_processor import process_pptx
from visual_processor import VisualAnalysis


class TestPptxProcessor(unittest.TestCase):

    def test_extracts_text_tables_and_slide_numbers(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.pptx"

            presentation = Presentation()
            blank = presentation.slide_layouts[6]

            slide1 = presentation.slides.add_slide(blank)

            text1 = slide1.shapes.add_textbox(
                Inches(1), Inches(1), Inches(5), Inches(1)
            )
            text1.text = "Registration deadline: November 12, 2026"

            table_shape = slide1.shapes.add_table(
                2, 2,
                Inches(1), Inches(2), Inches(5), Inches(1)
            )
            table_shape.table.cell(0, 0).text = "Process"
            table_shape.table.cell(0, 1).text = "Owner"
            table_shape.table.cell(1, 0).text = "Registration"
            table_shape.table.cell(1, 1).text = "Registrar"

            slide2 = presentation.slides.add_slide(blank)

            text2 = slide2.shapes.add_textbox(
                Inches(1), Inches(1), Inches(5), Inches(1)
            )
            text2.text = "Technology support reference: TECH-739"

            presentation.save(path)

            elements = process_pptx(path)

            self.assertEqual(
                [element.content_type for element in elements],
                ["text", "table", "text"],
            )
            self.assertEqual(
                [element.page_number for element in elements],
                [1, 1, 2],
            )
            self.assertEqual(
                [element.element_index for element in elements],
                [1, 2, 3],
            )
            self.assertIn("Registrar", elements[1].content)
            self.assertEqual(
                elements[1].extraction_method,
                "python_pptx_table",
            )

    def test_analyzes_embedded_image(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "image.pptx"

            presentation = Presentation()
            slide = presentation.slides.add_slide(
                presentation.slide_layouts[6]
            )

            image_buffer = BytesIO()
            Image.new("RGB", (120, 80), "white").save(
                image_buffer, format="PNG"
            )
            image_buffer.seek(0)

            slide.shapes.add_picture(
                image_buffer,
                Inches(1), Inches(1),
                width=Inches(3),
            )

            presentation.save(path)

            settings = type(
                "TestSettings",
                (),
                {"vlm_provider": "ollama"},
            )()

            analysis = VisualAnalysis(
                content_type="diagram",
                content="Submit Request -> Advisor Review -> Approval",
                model="test-vision-model",
            )

            with patch(
                "pptx_processor.analyze_visual",
                return_value=analysis,
            ) as mock_analyze:
                elements = process_pptx(path, settings=settings)

            self.assertEqual(len(elements), 1)
            self.assertEqual(elements[0].content_type, "diagram")
            self.assertEqual(elements[0].page_number, 1)
            self.assertEqual(
                elements[0].extraction_method,
                "ollama_vision",
            )
            self.assertEqual(
                elements[0].extraction_model,
                "test-vision-model",
            )
            mock_analyze.assert_called_once()

    def test_reuses_analysis_for_duplicate_images(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "duplicates.pptx"

            presentation = Presentation()
            blank = presentation.slide_layouts[6]

            image_buffer = BytesIO()
            Image.new("RGB", (120, 80), "white").save(
                image_buffer, format="PNG"
            )
            image_bytes = image_buffer.getvalue()

            for _ in range(2):
                slide = presentation.slides.add_slide(blank)
                slide.shapes.add_picture(
                    BytesIO(image_bytes),
                    Inches(1), Inches(1),
                    width=Inches(3),
                )

            presentation.save(path)

            settings = type(
                "TestSettings",
                (),
                {"vlm_provider": "ollama"},
            )()

            analysis = VisualAnalysis(
                content_type="image",
                content="A reusable informational image.",
                model="test-vision-model",
            )

            with patch(
                "pptx_processor.analyze_visual",
                return_value=analysis,
            ) as mock_analyze:
                elements = process_pptx(path, settings=settings)

            self.assertEqual(len(elements), 2)
            self.assertEqual(
                [element.page_number for element in elements],
                [1, 2],
            )
            mock_analyze.assert_called_once()


if __name__ == "__main__":
    unittest.main()
